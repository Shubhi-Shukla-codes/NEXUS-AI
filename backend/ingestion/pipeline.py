"""
NEXUS Document Ingestion & Chunking Pipeline
Coordinates parser dispatch, text normalization, page-level persistence,
structure-aware chunking, and SQLite relational synchronization.
"""

from pathlib import Path
from typing import Optional, Dict, Any, List
from backend.storage.db import get_db_connection
from backend.ingestion.cleaner import clean_text, remove_repetitive_headers_footers
from backend.ingestion.chunker import chunk_document_pages
from backend.ingestion.extractors.pdf_extractor import extract_pdf
from backend.ingestion.extractors.docx_extractor import extract_docx
from backend.ingestion.extractors.tabular_extractor import extract_xlsx
from backend.ingestion.extractors.text_extractor import extract_plaintext, extract_csv, extract_json
from backend.security.audit import log_audit_event


class PipelineError(Exception):
    """Raised when extraction or chunking fails."""
    pass


def extract_document_pages(file_path: Path, ext: str) -> List[Dict[str, Any]]:
    """Dispatch to correct format extractor based on file extension."""
    ext = ext.lower()
    if ext == ".pdf":
        return extract_pdf(file_path)
    elif ext == ".docx":
        return extract_docx(file_path)
    elif ext == ".xlsx":
        return extract_xlsx(file_path)
    elif ext == ".csv":
        return extract_csv(file_path)
    elif ext == ".json":
        return extract_json(file_path)
    elif ext in (".txt", ".md"):
        return extract_plaintext(file_path)
    else:
        raise PipelineError(f"Unsupported document format: {ext}")


def process_document(
    document_id: int,
    chunk_size: Optional[int] = None,
    chunk_overlap: Optional[int] = None,
    custom_path: Optional[Path] = None
) -> Dict[str, Any]:
    """
    Execute complete extraction, cleaning, and chunking pipeline for a document.
    Persists pages into 'document_pages' and chunks into 'document_chunks'.
    """
    with get_db_connection(custom_path) as conn:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, filename, vault_path, status, domain
            FROM documents WHERE id = ?
            """,
            (document_id,)
        )
        doc = cur.fetchone()
        if not doc:
            raise PipelineError(f"Document ID #{document_id} not found in repository.")

        vault_path = Path(doc["vault_path"])
        filename = doc["filename"]
        ext = vault_path.suffix.lower()

        # Update status to extracting
        cur.execute("UPDATE documents SET status = 'extracting' WHERE id = ?", (document_id,))

    if not vault_path.is_file():
        with get_db_connection(custom_path) as conn:
            conn.execute("UPDATE documents SET status = 'failed' WHERE id = ?", (document_id,))
        raise PipelineError(f"Physical vault file missing at: {vault_path}")

    try:
        # 1. Format-specific extraction
        raw_pages = extract_document_pages(vault_path, ext)
        if not raw_pages:
            raise PipelineError(f"No extractable text found in '{filename}'.")

        # 2. Text cleaning & normalization
        cleaned_pages = []
        for p in raw_pages:
            c_text = clean_text(p["text"])
            cleaned_pages.append({
                "page_number": p["page_number"],
                "text": c_text,
                "char_count": len(c_text)
            })

        # Strip recurring headers/footers
        final_pages = remove_repetitive_headers_footers(cleaned_pages)

        # Scanned page / diagram fallback: ensure visual pages have searchable provenance
        for p in final_pages:
            if len(p["text"].strip()) < 30:
                p["text"] = f"[VISUAL TECHNICAL ASSET: Page {p['page_number']} contains an authorized schematic, diagram, or tabular specification for '{filename}'.]"
                p["char_count"] = len(p["text"])

        # Multimodal diagram & table extraction for PDFs
        if ext == ".pdf":
            try:
                import fitz
                from backend.ingestion.multimodal import extract_images_from_pdf_page
                from backend.ingestion.table_parser import extract_tables_from_pdf_page, save_table_asset

                pdf_doc = fitz.open(vault_path)
                for page_num, pdf_page in enumerate(pdf_doc, start=1):
                    extract_images_from_pdf_page(
                        doc=pdf_doc,
                        page=pdf_page,
                        document_id=document_id,
                        page_number=page_num,
                        custom_path=custom_path
                    )
                    tabs = extract_tables_from_pdf_page(pdf_page)
                    for tab in tabs:
                        save_table_asset(
                            document_id=document_id,
                            page_number=page_num,
                            table_index=tab["table_index"],
                            headers=tab["headers"],
                            rows=[[r.get(h, "") for h in tab["headers"]] for r in tab["json_records"]],
                            custom_path=custom_path
                        )
                pdf_doc.close()
            except Exception as m_exc:
                pass
        elif ext == ".xlsx":
            try:
                import openpyxl
                from backend.ingestion.table_parser import save_table_asset
                wb = openpyxl.load_workbook(str(vault_path), data_only=True, read_only=True)
                for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
                    ws = wb[sheet_name]
                    rows = list(ws.iter_rows(values_only=True))
                    if rows and len(rows) > 1:
                        headers = [str(c or "").strip() for c in rows[0]]
                        data_rows = rows[1:]
                        save_table_asset(
                            document_id=document_id,
                            page_number=sheet_idx,
                            table_index=1,
                            headers=headers,
                            rows=data_rows,
                            custom_path=custom_path
                        )
                wb.close()
            except Exception as x_exc:
                pass

        # 3. Structure-aware chunking
        chunks = chunk_document_pages(
            document_id=document_id,
            pages=final_pages,
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap
        )

        if not chunks:
            raise PipelineError("Chunking resulted in 0 chunks.")

        # 4. Relational Persistence
        with get_db_connection(custom_path) as conn:
            cur = conn.cursor()

            # Insert/replace document pages
            cur.execute("DELETE FROM document_pages WHERE document_id = ?", (document_id,))
            for p in final_pages:
                cur.execute(
                    """
                    INSERT INTO document_pages (document_id, page_number, char_count)
                    VALUES (?, ?, ?)
                    """,
                    (document_id, p["page_number"], p["char_count"])
                )

            # Insert/replace document chunks
            cur.execute("DELETE FROM document_chunks WHERE document_id = ?", (document_id,))
            for c in chunks:
                cur.execute(
                    """
                    INSERT INTO document_chunks (
                        id, document_id, page_number, chunk_index,
                        content, char_length, token_count
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        c["id"],
                        c["document_id"],
                        c["page_number"],
                        c["chunk_index"],
                        c["content"],
                        c["char_length"],
                        c["token_count"]
                    )
                )

            # Update document chunk_count and transition status
            # If document was previously approved, retain approved; otherwise mark processed
            new_status = "approved" if doc["status"] == "approved" else "processed"
            cur.execute(
                """
                UPDATE documents
                SET chunk_count = ?, status = ?
                WHERE id = ?
                """,
                (len(chunks), new_status, document_id)
            )

        log_audit_event(
            event_type="DOC_EXTRACTED",
            actor_username="INGESTION_ENGINE",
            description=f"Extracted {len(final_pages)} pages and generated {len(chunks)} chunks for '{filename}'",
            metadata={"document_id": document_id, "pages": len(final_pages), "chunks": len(chunks)},
            custom_path=custom_path
        )

        return {
            "document_id": document_id,
            "filename": filename,
            "pages_count": len(final_pages),
            "chunks_count": len(chunks),
            "status": new_status
        }

    except Exception as exc:
        with get_db_connection(custom_path) as conn:
            conn.execute("UPDATE documents SET status = 'failed' WHERE id = ?", (document_id,))
        log_audit_event(
            event_type="DOC_EXTRACT_FAILED",
            actor_username="INGESTION_ENGINE",
            description=f"Extraction failed for document #{document_id} ('{filename}'): {exc}",
            metadata={"document_id": document_id, "error": str(exc)},
            custom_path=custom_path
        )
        raise PipelineError(f"Ingestion pipeline failed for document #{document_id}: {exc}")


def get_document_chunks(
    document_id: int,
    custom_path: Optional[Path] = None
) -> List[Dict[str, Any]]:
    """Retrieve all chunks belonging to a document."""
    with get_db_connection(custom_path) as conn:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT id, document_id, page_number, chunk_index,
                   content, char_length, token_count, created_at
            FROM document_chunks
            WHERE document_id = ?
            ORDER BY chunk_index ASC
            """,
            (document_id,)
        )
        return [dict(r) for r in cur.fetchall()]
