"""
NEXUS Excel Spreadsheet Extractor
Extracts multi-sheet tables and cell matrices using openpyxl.
"""

from pathlib import Path
from typing import List, Dict, Any
import openpyxl


def extract_xlsx(file_path: Path) -> List[Dict[str, Any]]:
    """
    Extract sheets and rows from an Excel spreadsheet.
    Each sheet is represented as a distinct logical page.
    """
    wb = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
    pages: List[Dict[str, Any]] = []

    for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            continue

        # Find header
        header = [str(c or "").strip() for c in rows[0]]
        text_lines = [f"Spreadsheet Sheet: '{sheet_name}'", f"Headers: {', '.join(header)}"]

        for r_idx, row in enumerate(rows[1:], start=2):
            non_empty = [f"{h}: {val}" for h, val in zip(header, row) if val is not None and str(val).strip()]
            if non_empty:
                text_lines.append(f"Row {r_idx}: " + " | ".join(non_empty))

        sheet_text = "\n".join(text_lines)
        pages.append({
            "page_number": sheet_idx,
            "text": sheet_text.strip(),
            "char_count": len(sheet_text.strip())
        })

    wb.close()
    return pages
